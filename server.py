"""Local dashboard with a shared 60-second Google Sheets cache."""
import io
import json
import threading
import time
import urllib.request
from datetime import date, datetime, timezone
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
SOURCES = [
    '1eh0xd3i_ZWCqP-xoH5h--K76jE4lE59BvHtVGZHfBfM',  # CD 961
    '1rxtJOqOvDuE_OGznvklIo8lzG5o5ActE',              # CD 969
    '1pqnr6JucXuP2h6cuBspwEb-0WSNUiQKR',              # CD 962
]
SHEETS = ['CAPACIDAD', 'ANTIGUEDAD', 'QUIEBRE', 'RECEPCION', 'PICKING', 'DESPACHO', 'RRHH', 'M. CONOC', 'M. DESCO']
SPANISH_MONTHS = {'ene': 1, 'feb': 2, 'mar': 3, 'abr': 4, 'may': 5, 'jun': 6, 'jul': 7, 'ago': 8, 'set': 9, 'sep': 9, 'oct': 10, 'nov': 11, 'dic': 12}
LOCK = threading.Lock()
CACHE = None
LAST_ATTEMPT = 0
LAST_ERROR = None

def normalize_date(value):
    """Keep spreadsheet dates consistent even when a sheet uses e.g. 21-Set."""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if not isinstance(value, str):
        return value
    parts = value.strip().lower().replace('/', '-').split('-')
    if len(parts) == 2 and parts[0].isdigit() and parts[1] in SPANISH_MONTHS:
        return date(datetime.now().year, SPANISH_MONTHS[parts[1]], int(parts[0])).isoformat()
    return value

def read_sources():
    data = {name: [] for name in SHEETS}
    for source in SOURCES:
        request = urllib.request.Request(f'https://docs.google.com/spreadsheets/d/{source}/export?format=xlsx', headers={'User-Agent': 'OperationalDashboard/1.0'})
        with urllib.request.urlopen(request, timeout=25) as response:
            workbook = load_workbook(io.BytesIO(response.read()), read_only=True, data_only=True)
        try:
            for name in SHEETS:
                rows = workbook[name].iter_rows(values_only=True)
                headers = [str(v).strip() if v is not None else '' for v in next(rows)]
                if 'COD_CD' not in headers:
                    raise ValueError(f'Encabezados no válidos: {name}')
                for values in rows:
                    row = {key: (normalize_date(value) if key == 'FECHA' else (value.isoformat() if isinstance(value, (date, datetime)) else value))
                           for key, value in zip(headers, values) if key}
                    if any(value is not None for value in row.values()):
                        data[name].append(row)
        finally:
            workbook.close()
    return {'data': data, 'updatedAt': datetime.now(timezone.utc).isoformat()}

class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def do_GET(self):
        global CACHE, LAST_ATTEMPT, LAST_ERROR
        if self.path.split('?')[0] != '/api/data':
            return super().do_GET()
        with LOCK:
            if time.monotonic() - LAST_ATTEMPT >= 60 or LAST_ATTEMPT == 0:
                LAST_ATTEMPT = time.monotonic()
                try:
                    CACHE = read_sources()
                    LAST_ERROR = None
                except Exception:
                    LAST_ERROR = 'No se pudo consultar Google Sheets. Se conservan los últimos datos.'
            payload = dict(CACHE or {})
            payload['stale'] = LAST_ERROR is not None
            if LAST_ERROR:
                payload['error'] = LAST_ERROR
        body = json.dumps(payload, ensure_ascii=False, allow_nan=False).encode('utf-8')
        self.send_response(200 if CACHE else 503)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Cache-Control', 'no-store')
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

if __name__ == '__main__':
    ThreadingHTTPServer(('127.0.0.1', 8765), Handler).serve_forever()
